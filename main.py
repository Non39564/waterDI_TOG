from flask import Flask, make_response, jsonify, request, Response, flash, redirect, send_file
from flask_login import LoginManager
import flask_login, flask, generate, json, docx, os, base64
from sendemail import *
from flask import render_template as real_render_template
from functools import partial
from db import *
from document import *
from datetime import datetime
from pymongo import MongoClient
from docx import *
import xlsxwriter

app = Flask(__name__)
app.secret_key = 'waterdishow'
app.config['UPLOAD_FOLDER'] = './static/signature/'
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}

error = showerror()
user_recept = show_recept()
passive = "False"
MM = "False"
render_template = partial(real_render_template, error=error, user_recept=user_recept, passive = passive, MM = MM)

def allowed_file(filename):
    return '.' in filename and \
        filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

dataexport = 0

@app.route('/postdata', methods=["GET", "POST"])
def postdata():
    if request.method == 'POST':
        print("post")
        data = request.json
        now = datetime.now()
        sec = now.strftime("%S")
        json_object = json.dumps(data)
        with open("data.json", "w") as outfile:
            outfile.write(json_object)
        
        # if int(sec) % 10 == 0 :
        #     client = MongoClient('localhost',27017)
        #     db = client.Water_di
        #     tb = db['report']
        #     tb.insert_many(data)
        #     print('insert to mogodb')
        #     dataexport = 0


        return "postdata succress"

@app.route('/dataapi', methods=["GET", "POST"])
def dataapi():
    js = open('data.json')
    data = json.load(js) 
    response = make_response(json.dumps(data))
    response.content_type = 'application/json'
    return jsonify(data)

########################### Function Login ###########################
login_manager = LoginManager()
login_manager.init_app(app)

users = users()
permission = users_permission()

class User(flask_login.UserMixin):
    pass


@login_manager.user_loader
def user_loader(email):
    if email not in users:
        return

    user = User()
    user.id = email
    return user

@login_manager.request_loader
def request_loader(request):
    email = request.form.get('email')
    if email not in users:
        return

    user = User()
    user.id = email
    return user

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        if flask.request.form['check'] == "login":
            username = flask.request.form['username']
            if username in users and flask.request.form['password'] == users[username]['password']:
                user = User()
                user.id = username
                flask_login.login_user(user)
                if 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
                    passive = 'True'
                    MM = 'False'
                elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
                    passive = 'False'
                    MM = 'True'
                else:
                    passive = 'False'
                    MM = 'False'
                return flask.redirect(flask.url_for('index', passive = passive, MM = MM))
            
            errorpassword = "True"
            return render_template('index.html', sweetalert=errorpassword)
        if flask.request.form['check'] == "register":
            Name = flask.request.form['registerName']
            username = flask.request.form['registerUsername']
            password = flask.request.form['registerPassword']
            password_again = flask.request.form['registerRepeatPassword']
            department = flask.request.form['DP']
            if 'file' not in request.files:
                flash('No file part')
                return redirect(request.url)
            Signature = request.files['file']
            if Signature.filename == '':
                flash('No selected file')
                return redirect(request.url)
            if password == password_again:
                if Signature and allowed_file(Signature.filename):
                    Signature.save(os.path.join(app.config['UPLOAD_FOLDER'], username+'.png'))
                    file = base64.b64encode(open(f'./static/signature/{username}.png','rb').read())
                    insert_Pre_User(username, Name, password, file, department)
                    waitforrecept = "True"
                    return render_template('index.html', waitforrecept=waitforrecept)
            else:
                notmatchpassword = "True"
                return render_template('index.html', notmatchpassword=notmatchpassword)

@app.route('/addmachine', methods=['GET', 'POST'])
@flask_login.login_required
def addmachine():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    if request.method == 'POST':
        if flask.request.form['add'] == "machine":
            id = flask.request.form['MachineId']
            ip = flask.request.form['MachineIp']
            port = flask.request.form['MachinePort']
            station = flask.request.form['inputOpMac']
            phase = flask.request.form['inputPhaseMac']
            Add_Machine(id, ip, port, station, phase)
            return flask.redirect(flask.url_for('addmachine'))
        if flask.request.form['add'] == "device":
            Machine = flask.request.form['inputMachineID']
            Site = flask.request.form['inputSite']
            Slot_Temp = flask.request.form['Slot_temp']
            Slot_Water= flask.request.form['Water']
            Add_Device(Machine, Site, Slot_Temp, Slot_Water)
            return flask.redirect(flask.url_for('addmachine'))
        if flask.request.form['add'] == "phase":
            OPadd = flask.request.form['OPadd']
            Phaseadd = flask.request.form['Phaseadd']
            Siteadd = flask.request.form['Siteadd']
            add_phase(OPadd, Phaseadd, Siteadd)
            return flask.redirect(flask.url_for('addmachine'))
        
    station = get_dropdown_values()
    machine_data = get_dropdown_values_machine()
    slot = dynamic_slot()
    edit_delete = show_machine()
    site = get_values_site()
    
    return render_template('add.html', station = station, machine_data = machine_data, slot = slot, edit_delete = edit_delete, site = site, passive = passive, MM = MM)

@app.route('/delete', methods=['GET', 'POST'])
@flask_login.login_required
def delete():
    Site = request.args.get('Site')
    delete_device(Site)
    return flask.redirect(flask.url_for('addmachine'))

@app.route('/edit', methods=['GET', 'POST'])
@flask_login.login_required
def edit():
    Site = request.args.get('edit')
    edit = edit_machine_device(Site)
    station = get_dropdown_values()
    slot = dynamic_slot()
    return render_template('edit_device.html', edit = edit, station = station, slot = slot)

@app.route('/blank')
@flask_login.login_required
def blank():
    return render_template('blank.html')

@app.route('/edit_value', methods=['GET', 'POST'])
@flask_login.login_required
def edit_value():
    if request.method == 'POST':
        result = request.form.to_dict()
        values = result.values()
        key = list(result.keys())
        values_Site = list(values)
        Site = key[0]
        Ip = request.form.get("Ip"+values_Site[0])
        Port = request.form.get("Port"+values_Site[0])
        Station = request.form.get("inputOP")
        Phase = request.form.get("Phase"+values_Site[0])
        Slot_Water = request.form.get("Slot_Water"+values_Site[0])
        Slot_Temp = request.form.get("Slot_Temp"+values_Site[0])
        if Slot_Water is None and Slot_Temp is None:
            edit = edit_machine_device(Site)
            Slot_Water = edit[0]['Slot_Water']
            Slot_Temp = edit[0]['Slot_Temp']
        print(Site, Ip, Port, Station, Phase, Slot_Water, Slot_Temp)
        edit_device(Ip, Port, Station, Phase, Slot_Water, Slot_Temp, Site)
    return flask.redirect(flask.url_for('blank'))
        
@app.route('/logout')
def logout():
    flask_login.logout_user()
    return flask.redirect(flask.url_for('index'))

@login_manager.unauthorized_handler
def unauthorized_handler():
    return 'Unauthorized', 401
    
@app.route('/setup')
@flask_login.login_required
def setup():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    setup = setupMachine()
    show_machine = machine_DropDown_setup()
    return render_template('setup.html', setup = setup, show_machine = show_machine,passive = passive, MM = MM)

@app.route('/find_setup')
@flask_login.login_required
def find_setup():
    selectionMachine = request.args.get('selectionMachine')
    setup = show_site_machine(selectionMachine)
    show_machine = machine_DropDown_setup()
    return render_template('setup.html', setup = setup, show_machine = show_machine, selectionMachine=selectionMachine)

@app.route('/confirm_reject', methods=['GET', 'POST'])
@flask_login.login_required
def confirm_reject():
    if request.args.get('confirm'):
        user_confirm = request.args.get('confirm')
        confirm_User(user_confirm)
        return flask.redirect(flask.url_for('index'))
    if request.args.get('reject'):
        user_reject = request.args.get('reject')
        reject_User(user_reject)
        return flask.redirect(flask.url_for('index'))
    return flask.redirect(flask.url_for('index'))

@app.route('/change_setup' , methods=['GET', 'POST'])
@flask_login.login_required
def change_setup():
    setup = setupMachine()
    name_site = []
    for i in setup:
        name_site.append(i['Site'])
    for j in range(len(name_site)):
        site = request.args.get(f'{name_site[j]}')
        low_water = request.args.get(f'Low_Water_{name_site[j]}')
        high_water = request.args.get(f'High_Water_{name_site[j]}')
        plus_water = request.args.get(f'Plus_Water_{name_site[j]}')
        minus_water = request.args.get(f'Minus_Water_{name_site[j]}')
        plus_temp = request.args.get(f'Plus_Temp_{name_site[j]}')
        minus_temp = request.args.get(f'Minus_Temp_{name_site[j]}')
        update_setup(site, low_water, high_water, plus_water, minus_water, plus_temp, minus_temp)
    return flask.redirect(flask.url_for('setup'))

    
@app.route('/sendfileit', methods=['POST', 'GET'])
@flask_login.login_required
def sendfileit():
    document_moc_or_auto('./static/document/ใบขอใช้บริการระบบงานคอมพิวเตอร์ - สำเนา.docx')
    template = './static/document/test.docx'
    signature = f'./static/signature/{flask_login.current_user.get_id()}.png'
    invoice = {
        'do': flask.request.form['cando'],
        'detail' : flask.request.form['system-effect'],
        'date_success': flask.request.form['date_it_finish'],
        'date_success_end': flask.request.form['date_it_deliver'],
        'date_update': flask.request.form['date_it_getWork'],
        'name': flask.request.form['nameit'],
        'date_sent': flask.request.form['date_it_now'],
        'all_day': flask.request.form['system-effect'],
    }
    document = generate.from_template(template, signature, invoice)
    document.seek(0)
    document.save('./static/document/ไอทีรับทราบ.docx')
    
    return send_file(
        document, mimetype='application/vnd.openxmlformats-'
        'officedocument.wordprocessingml.document', as_attachment=True,
        attachment_filename='ไอทีรับทราบ.docx')

@app.route('/confirm_request')
@flask_login.login_required
def confirm_request():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    KeyWork = request.args.get('KeyWork')
    result = detail_timeline(KeyWork)
    return render_template('confirm-request.html', result = result, passive = passive, MM = MM)

@app.route('/document')
@flask_login.login_required
def document():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    if 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        Access = "True"
        SelectDoc = SelectDoc_IT()
        return render_template('request-document.html', SelectDoc = SelectDoc, passive = passive, MM = MM)
    else:
        Access = "False"
        return render_template('document.html', Access=Access, passive = passive, MM = MM)
    
@app.route('/document_it')
@flask_login.login_required
def document_it():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    if 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        Work_id = request.args.get('KeyWork')
        NameWork = request.args.get('NameWork')
        Access = "True"
        return render_template('document.html', Access=Access, Work_id = Work_id, Work = NameWork,passive = passive, MM = MM)

@app.route('/update_doc_it', methods=['GET', 'POST'])
@flask_login.login_required
def update_doc_it():
    if request.method == 'POST':
        result = request.form.to_dict()
        values = result.values()
        values_data = list(values)
        Username = flask_login.current_user.get_id()
        Cando = values_data[0]
        Effect = values_data[1]
        date_finish = values_data[3]
        date_deliver = values_data[4]
        date_getWork = values_data[5]
        date = values_data[5]
        KeyWork = values_data[6]
        Sum_date = (datetime.strptime(values_data[3], '%Y-%m-%d') - datetime.strptime(values_data[4], '%Y-%m-%d'))
        Detail = values_data[2]
        Update_Doc_IT(Username, Cando, Effect, Detail, date_finish, date_deliver, date_getWork, Sum_date, KeyWork)
        return flask.redirect(flask.url_for('blank'))
    
@app.route('/update_status_doc', methods=['GET', 'POST'])
@flask_login.login_required
def update_status_doc():
    if request.method == 'POST':
        result = request.form.to_dict()
        values = result.values()
        values_data = list(values)
        Status = values_data[0]
        KeyWork = values_data[1]
        if 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
            Access = "True"
            SelectDoc = SelectDoc_IT()
            update_state_doc(Status, KeyWork)
            return flask.redirect(flask.url_for('document', SelectDoc = SelectDoc, Access = Access))

@app.route('/edit_maintain', methods=['GET', 'POST'])
@flask_login.login_required
def edit_maintain():
    Site = request.args.get('edit_maintain')
    if 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        edit = edit_machine_device(Site)
        longname = find_user_data(flask_login.current_user.get_id())
        now = datetime.now()
        time = now.strftime("%H:%M:%S")
        date = now.strftime("%Y-%m-%d")
        return render_template('popup-acknow.html', longname = longname, edit = edit, Date = date, Time = time)
    return render_template('popup-acknow.html')

@app.route('/edit_maintain_value', methods=['GET', 'POST'])
@flask_login.login_required
def edit_maintain_value():
    if request.method == 'POST':
        result = request.form.to_dict()
        values = result.values()
        values_Site = list(values)
        Site = values_Site[0]
        State = values_Site[7]
        Date = values_Site[5]
        Time = values_Site[6]
        update_process(flask_login.current_user.get_id(), State, Site)
        insert_process(flask_login.current_user.get_id(), State, Site, Date, Time)
    return flask.redirect(flask.url_for('blank'))

@app.route('/popup-acKnow')
@flask_login.login_required
def popupAcknow():
    return render_template('popup-acknow.html')

@app.route('/insertDoc', methods=['POST', 'GET'])
@flask_login.login_required
def insertDoc():
    Username = flask_login.current_user.get_id()
    date = flask.request.form['datefornow']
    time = flask.request.form['timefornow']
    place = flask.request.form['place']
    service = flask.request.form['service']
    reason = flask.request.form['reason']
    system_now = flask.request.form['system_now']
    detail = flask.request.form['detail']
    date_end = flask.request.form['date_end']
    note = flask.request.form['note']
    insertDocument(Username, date, time, place, service, reason, system_now, detail, date_end, note)
    for i in countdataDoc():
        KeyWork = i["alldata"]
    insertDocumentit(KeyWork)
    
    userdata = find_user_data(flask_login.current_user.get_id())
    document_moc_or_auto('./static/document/ใบขอใช้บริการระบบงานคอมพิวเตอร์ - สำเนา.docx')
    template = './static/document/test.docx'
    invoice = {
        'name': userdata[0]["Name"],
        'position' : userdata[0]["position"],
        'department' : userdata[0]["Department"],
        'phone': userdata[0]["Phone"],
        'email': userdata[0]["Email"],
        'part': userdata[0]["part"],
        'date': flask.request.form['datefornow'],
        'time': flask.request.form['timefornow'],
        'place': flask.request.form['place'],
        'service': flask.request.form['service'],
        'reason': flask.request.form['reason'],
        'system_now': flask.request.form['system_now'],
        'detail': flask.request.form['detail'],
        'date_end': flask.request.form['date_end'],
        'note': flask.request.form['note'],
        'signature' :userdata[0]["Name"]
    }
    document = generate.from_template(template, invoice)
    document.seek(0)
    doc = docx.Document(document)
    doc.save('./static/document/แบบคำขอใช้งานส่งไอที.docx')
    email(userdata[0]["Name"], flask.request.form['service'], str(flask.request.form['date_end']))
    return flask.redirect(flask.url_for('requestTimeline'))

@app.route('/request-timeline')
@flask_login.login_required
def requestTimeline():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    Timeline = Select_Timeline()
    return render_template('request-timeline.html', Select_Timeline = Timeline, passive = passive, MM = MM)

@app.route('/view-process')
@flask_login.login_required
def viewProcess():
    maintain = find_Maintain()
    if flask_login.current_user.is_anonymous:
        passive = "False"
        Access = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        Access = "True"
        passive = "True"
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
        Access = "False"
    else :
        Access = "False"
        passive = "False"
        MM = 'False'
    return render_template('view-process.html', maintain = maintain, Access = Access, passive = passive, MM = MM) 
    
        
############################################################## End Function Login #################################################

@app.route('/')
def index():
    station = get_dropdown_values()
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    return render_template('index.html', station = station, passive = passive, MM = MM)

@app.route("/data_line", methods=["POST","GET"])
def data_line():
    
    if request.method == 'POST':
        startdate = request.form['startdate']
        enddate =  request.form['enddate']
        draw = request.form['draw']
        row = int(request.form['start'])
        rowperpage = int(request.form['length'])
        searchValue = request.form["search[value]"]
        
        report = di_report_now(startdate, enddate)
        totalRecords = len(report)
        
        likeString = "%" + searchValue +"%"
        filter = di_report_filter_table(startdate, enddate, likeString)
        print(likeString)
        totalRecordwithFilter = len(filter)

        if searchValue == '':
            report = di_report_limit(startdate, enddate, row, rowperpage)
        else:
            report = di_report_filter_table_limit(startdate, enddate, likeString,row, rowperpage)
        data = []
        for d in report:
            data.append({
                'Status': d['State'],
                'Phase': d['Phase'],
                'Site': d['Site'],
                'DIWater': d['Water'],
                'Date': d['Date'],
                'Time': str(d['Time']),
                'Temp':d['Temp']
            })
        response = {
                'draw': draw,
                'iTotalRecords': totalRecordwithFilter,
                'iTotalDisplayRecords': totalRecords,
                'aaData': data,
            }
        return jsonify(response)

@app.route('/line', methods=['GET', 'POST'])
def line():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    if request.method == 'POST':
        date = request.values["pick_date_chart"]
        report = di_report_custom(date)
        table = report_line_month(date)
        column = somlinecolumn()
        custom = "True"
        return render_template('line.html', report = report, table = table, column = column, custom = custom, passive = passive, MM = MM)
    table = reportsomline()
    column = somlinecolumn()
    now = "True"
    return render_template('line.html', table = table, column = column, now = now, passive = passive, MM = MM)

@app.route('/report')
def report():
    return render_template('report.html')


@app.route("/data_alert", methods=["POST","GET"])
def data_alert():
    if request.method == "POST":
        startdate = request.form['startdate']
        enddate =  request.form['enddate']
        draw = request.form['draw'] 
        row = int(request.form['start'])
        rowperpage = int(request.form['length'])
        searchValue = request.form["search[value]"]
        
        report = error_report_find(startdate, enddate)
        totalRecords = len(report)
        
        likeString = "%" + searchValue +"%"
        filter = filter_table(startdate, enddate, likeString)
        totalRecordwithFilter = len(filter)
        
        if searchValue == '':
            report = error_report_limit(startdate, enddate, row, rowperpage)
        else:
            report = filter_table_limit(startdate,enddate,likeString,row, rowperpage)
          
        data = []
        for d in report:
            data.append({
                'Station': d['Station'],
                'Phase': d['Phase'],
                'Site': d['Site'],
                'Detail': d['Detail'],
                'Date': d['Date'],
                'Time': str(d['Time']),
                'Date_Time': (datetime.combine(d['Date'],(datetime.min + d['Time']).time()))
            })
        response = {
                'draw': draw,
                'iTotalRecords': totalRecordwithFilter,
                'iTotalDisplayRecords': totalRecords,
                'aaData': data,
            }
        return jsonify(response)

@app.route('/alert')
def alert():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    totalRecords = len(error_report())
    return render_template('alert.html', totalRecords = totalRecords,passive = passive, MM = MM)

@app.route('/export_alert', methods=["POST","GET"])
def export_alert():
    startdate = request.form['startdate']
    enddate = request.form['enddate']
    print(startdate)
    print(enddate)
    Site = get_error_name()
    report = error_report_find(startdate, enddate)
    workbook   = xlsxwriter.Workbook('Alert_report.xlsx')
    num = 1
    for i in range(len(Site)):
        worksheet = workbook.add_worksheet(Site[i])
        worksheet.write(0, 0, 'Station') 
        worksheet.write(0, 1, 'Phase') 
        worksheet.write(0, 2, 'Site') 
        worksheet.write(0, 3, 'Detail') 
        worksheet.write(0, 4, 'Date') 
        worksheet.write(0, 5, 'Time') 
        for j in range(len(report)):
            if report[j]['Site'] == Site[i]:
                worksheet.write(num, 0, report[j]['Station']) 
                worksheet.write(num, 1, report[j]['Phase']) 
                worksheet.write(num, 2, report[j]['Site']) 
                worksheet.write(num, 3, report[j]['Detail']) 
                worksheet.write(num, 4, str(report[j]['Date'])) 
                worksheet.write(num, 5, str(report[j]['Time'])) 
                num += 1
        num = 1
    workbook.close()
    now = datetime.now()
    date = now.strftime("%Y%m%d")
    return send_file('Alert_report.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'Alert_report_{date}.xlsx')
    
@app.route('/Table-of-Deionized-water', methods=["POST","GET"])
def Table_of_Deionized_water():
    startdate = request.form['startdate']
    enddate = request.form['enddate']
    Site = report_Site_Name()
    report = di_report_now(startdate, enddate)
    workbook   = xlsxwriter.Workbook('Table of Deionized water.xlsx')
    num = 1
    for i in range(len(Site)):
        worksheet = workbook.add_worksheet(Site[i])
        worksheet.write(0, 0, 'Date') 
        worksheet.write(0, 1, 'Time') 
        worksheet.write(0, 2, 'Phase') 
        worksheet.write(0, 3, 'Site') 
        worksheet.write(0, 4, 'Status') 
        worksheet.write(0, 5, 'Water')
        worksheet.write(0, 6, 'Temp') 
        for j in range(len(report)):
            if report[j]['Site'] == Site[i]:
                worksheet.write(num, 0, str(report[j]['Date'])) 
                worksheet.write(num, 1, str(report[j]['Time'])) 
                worksheet.write(num, 2, report[j]['Phase']) 
                worksheet.write(num, 3, report[j]['Site']) 
                worksheet.write(num, 4, report[j]['State']) 
                worksheet.write(num, 5, report[j]['Water']) 
                worksheet.write(num, 6, report[j]['Temp'])
                num += 1
        num = 1
    workbook.close()
    now = datetime.now()
    date = now.strftime("%Y%m%d")
    return send_file('Table of Deionized water.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'Table of Deionized water_{date}.xlsx')

@app.route('/export_trend', methods=["POST","GET"])
def export_trend():
    phase = request.args.get('phase')
    startdate = request.args.get('startdate')
    enddate = request.args.get('enddate')
    Site = get_trendDI_Site(phase)
    if phase == '4':
        report = trend_DI_P4(startdate,enddate)
    elif phase == '5':
        report = trend_DI_P5(startdate,enddate)
    elif phase == '9':
        report = trend_DI_P9(startdate,enddate)
    workbook   = xlsxwriter.Workbook(f'Trend_Water Di Phase {phase}.xlsx')
    num = 1
    print(Site)
    for i in range(len(Site)):
        worksheet = workbook.add_worksheet(Site[i])
        worksheet.write(0, 0, 'Date') 
        worksheet.write(0, 1, 'Time') 
        worksheet.write(0, 2, 'Site') 
        worksheet.write(0, 3, 'Water') 
        worksheet.write(0, 4, 'Temp') 
        for j in range(len(report)):
            if report[j]['Site'] == Site[i]:
                worksheet.write(num, 0, str(report[j]['Date']))
                worksheet.write(num, 1, str(report[j]['Time']))
                worksheet.write(num, 2, report[j]['Site'])
                worksheet.write(num, 3, report[j]['Water'])
                worksheet.write(num, 4, report[j]['Temp'])
                num += 1
        num = 0
    workbook.close()
    now = datetime.now()
    date = now.strftime("%Y%m%d")
    return send_file(f'Trend_Water Di Phase {phase}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        as_attachment=True, download_name=f'Trend_Water Di Phase {phase}_{date}.xlsx')     
    

@app.route('/trendDi')
def trendDi():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    startdate = ''
    enddate = ''
    columnforP4 = columnP4()
    tableforP4 = tableP4()
    columnforP9 = columnP9()
    tableforP9 = tableP9()
    columnforP5 = columnP5()
    tableforP5 = tableP5()
    RecordsP4 = len(trend_DI_P4(startdate,enddate))
    RecordsP5 = len(trend_DI_P5(startdate,enddate))
    RecordsP9 = len(trend_DI_P9(startdate,enddate))
    totalRecords = RecordsP4 + RecordsP5 + RecordsP9
    return render_template('TrendDiwater.html',totalRecords = totalRecords, columnP4 = columnforP4, columnP5 = columnforP5, columnP9 = columnforP9,
                           tableP4 = tableforP4, tableP5 = tableforP5, tableP9 = tableforP9, passive = passive, MM = MM)

@app.route("/data_di_P4", methods=["POST","GET"])
def data_di_P4():
    
    if request.method == 'POST':
        startdate = request.form['startdate']
        enddate =  request.form['enddate']
        draw = request.form['draw'] 
        row = int(request.form['start'])
        rowperpage = int(request.form['length'])
        searchValue = request.form["search[value]"]
        
        report = trend_DI_P4(startdate, enddate)
        totalRecords = len(report)
        
        likeString = "%" + searchValue +"%"
        filter = trend_DI_P4_filter_table(startdate, enddate, likeString)
        totalRecordwithFilter = len(filter)

        if searchValue == '':
            report = trend_DI_P4_limit(startdate, enddate, row, rowperpage)
        else:
            report = trend_DI_P4_filter_table_limit(startdate, enddate, likeString, row, rowperpage)
        data = []
        for d in report:
            data.append({
                'Phase': d['Phase'],
                'Site': d['Site'],
                'Water': d['Water'],
                'Date': d['Date'],
                'Time': str(d['Time']),
                'Temp':d['Temp']
            })
        response = {
                'draw': draw,
                'iTotalRecords': totalRecordwithFilter,
                'iTotalDisplayRecords': totalRecords,
                'aaData': data,
                }
        return jsonify(response)
    
@app.route("/data_di_P5", methods=["POST","GET"])
def data_di_P5():
    if request.method == 'POST':
        startdate = request.form['startdate']
        enddate =  request.form['enddate']
        draw = request.form['draw'] 
        row = int(request.form['start'])
        rowperpage = int(request.form['length'])
        searchValue = request.form["search[value]"]
        
        report = trend_DI_P5(startdate,enddate)
        totalRecords = len(report)
        
        likeString = "%" + searchValue +"%"
        filter = trend_DI_P5_filter_table(startdate,enddate,likeString)
        totalRecordwithFilter = len(filter)

        if searchValue == '':
            report = trend_DI_P5_limit(startdate,enddate,row, rowperpage)
        else:
            report = trend_DI_P5_filter_table_limit(startdate,enddate,likeString,row, rowperpage)
        data = []
        for d in report:
            data.append({
                'Phase': d['Phase'],
                'Site': d['Site'],
                'Water': d['Water'],
                'Date': d['Date'],
                'Time': str(d['Time']),
                'Temp':d['Temp']
            })
        response = {
                'draw': draw,
                'iTotalRecords': totalRecordwithFilter,
                'iTotalDisplayRecords': totalRecords,
                'aaData': data,
                }
        return jsonify(response)
    
@app.route("/data_di_P9", methods=["POST","GET"])
def data_di_P9():
    
    if request.method == 'POST':
        startdate = request.form['startdate']
        enddate =  request.form['enddate']
        draw = request.form['draw'] 
        row = int(request.form['start'])
        rowperpage = int(request.form['length'])
        searchValue = request.form["search[value]"]
        
        report = trend_DI_P9(startdate,enddate)
        totalRecords = len(report)
        
        likeString = "%" + searchValue +"%"
        filter = trend_DI_P9_filter_table(startdate,enddate,likeString)
        totalRecordwithFilter = len(filter)

        if searchValue == '':
            report = trend_DI_P9_limit(startdate,enddate,row, rowperpage)
        else:
            report = trend_DI_P9_filter_table_limit(startdate,enddate,likeString,row, rowperpage)
        data = []
        for d in report:
            data.append({
                'Phase': d['Phase'],
                'Site': d['Site'],
                'Water': d['Water'],
                'Date': d['Date'],
                'Time': str(d['Time']),
                'Temp':d['Temp']
            })
        response = {
                'draw': draw,
                'iTotalRecords': totalRecordwithFilter,
                'iTotalDisplayRecords': totalRecords,
                'aaData': data,
                }
        return jsonify(response)

@app.route('/trendDi-chart')
def trendDiChart():
    if flask_login.current_user.is_anonymous:
        passive = "False"
        MM = 'False'
    elif 3 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'True'
        MM = 'False'
    elif 1 == permission[flask_login.current_user.get_id()]['DepartmentID']:
        passive = 'False'
        MM = 'True'
    else:
        passive = 'False'
        MM = 'False'
    startdate = ''
    enddate = ''
    P4 = trend_DI_P4(startdate, enddate)
    P5 = trend_DI_P5(startdate, enddate)
    P9 = trend_DI_P9(startdate, enddate)
    return render_template('trend-di-chart.html', P4 = P4, P5 = P5, P9 = P9, passive= passive, MM = MM)

@app.route('/status-carbon-resin')
def statusCR():
    return render_template('status-carbon-resin.html')   

@app.route('/popup-process')
def popupProcess():
    Site = request.args.get('show')
    process = show_process(Site)
    edit = edit_machine_device(Site)
    return render_template('popup-process-status.html', process = process, edit = edit)

@app.route("/generator")
def generator():
    return render_template("index_generator.html")

@app.route('/postdata_generator', methods=["GET", "POST"])
def postdata_generator():
    if request.method == 'POST':
        print("post")
        data = request.json
        json_object = json.dumps(data)
        with open("data_elec.json", "w") as outfile:
            outfile.write(json_object)
        return "postdata succress"

@app.route('/dataapi_generator', methods=["GET", "POST"])
def dataapi_generator():
    js = open('data_elec.json')
    data_json = json.load(js) 
    response = make_response(json.dumps(data_json))
    response.content_type = 'application/json'
    return jsonify(data_json)


if __name__ == "__main__":
    app.config['SERVER_NAME'] = "10.3.9.156:80"
    app.run(debug=True)